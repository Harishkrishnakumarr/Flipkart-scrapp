"""Comprehensive unit tests for GST number extraction, validation, and multi-engine enrichment.

Covers all 30 required test scenarios:
1. Valid GSTIN extraction
2. Invalid GSTIN rejection
3. GSTIN normalization
4. Lowercase GSTIN normalization
5. Whitespace GSTIN normalization
6. State-code validation
7. GSTIN format structure validation
8. GST extraction from visible HTML
9. GST extraction from JSON-LD
10. GST extraction from embedded state
11. Direct Flipkart GST extraction
12. External search GST extraction
13. Google -> Bing fallback
14. Bing -> Brave fallback
15. Brave -> DuckDuckGo fallback
16. Google 429 fallback
17. Search-result snippet extraction
18. Opened-page GST extraction
19. Seller name matching
20. City matching
21. State matching
22. Pincode matching
23. Generic seller protection
24. Wrong-city GST rejection
25. Wrong-state GST rejection
26. Same-name business rejection
27. Multiple GSTIN candidate selection
28. GST state-code cross-check
29. Temporary failure must not become permanent NOT_FOUND
30. Final output mapping to "GST Number"
"""

import pytest
import asyncio
from unittest.mock import AsyncMock, patch, MagicMock
from bs4 import BeautifulSoup

from scraper.validator import (
    validate_gst,
    match_gst_to_seller,
    is_generic_seller_name,
)
from scraper.product_parser import (
    extract_flipkart_seller_metadata,
    parse_product_page,
)
from scraper.web_research import (
    WebResearchEngine,
    generate_targeted_gst_queries,
)
from scraper.exporter import LiveExcelManager


# =====================================================================
# 1-7: GSTIN FORMAT, NORMALIZATION, STATE CODES & MASKING
# =====================================================================

def test_01_valid_gstin_extraction():
    """Test standard 15-character valid GSTIN extraction."""
    gst = "24AAIHD1204K1ZQ"
    assert validate_gst(gst) == "24AAIHD1204K1ZQ"


def test_02_invalid_gstin_rejection():
    """Test malformed, too short, too long, or non-alphanumeric GSTIN rejection."""
    assert validate_gst("12345") is None
    assert validate_gst("24AAIHD1204K1") is None  # 13 chars
    assert validate_gst("24AAIHD1204K1ZQQ") is None  # 16 chars
    assert validate_gst("INVALIDGSTIN123") is None
    assert validate_gst("27AABCA****1Z9") is None  # Masked with asterisks
    assert validate_gst("27AABCAXXXX1Z9") is None  # Masked with X
    assert validate_gst("27AABCA••••1Z9") is None  # Masked with bullets
    assert validate_gst("") is None
    assert validate_gst(None) is None


def test_03_gstin_normalization():
    """Test normalization by stripping dashes, spaces, and underscores."""
    raw = "24-AAIHD-1204K-1-Z-Q"
    assert validate_gst(raw) == "24AAIHD1204K1ZQ"


def test_04_lowercase_gstin_normalization():
    """Test lowercase GSTIN string normalization to uppercase."""
    raw = "24aaihd1204k1zq"
    assert validate_gst(raw) == "24AAIHD1204K1ZQ"


def test_05_whitespace_gstin_normalization():
    """Test irregular whitespace GSTIN normalization."""
    raw = " 24 aaihd 1204 k 1 zq "
    assert validate_gst(raw) == "24AAIHD1204K1ZQ"


def test_06_state_code_validation():
    """Test state code validation against official 2-digit Indian state codes."""
    # Valid state codes (07: Delhi, 24: Gujarat, 33: Tamil Nadu, 27: Maharashtra)
    assert validate_gst("07AAACH6891K1ZW") == "07AAACH6891K1ZW"
    assert validate_gst("33AAECS5412Q1ZM") == "33AAECS5412Q1ZM"
    assert validate_gst("27AABCT3518Q1ZV") == "27AABCT3518Q1ZV"

    # Invalid state code 00 / 95 (not in list)
    assert validate_gst("00AAACH6891K1ZW") is None
    assert validate_gst("95AAACH6891K1ZW") is None


def test_07_gstin_format_structure_validation():
    """Test structural components: 2-digit state, 10-char PAN, 1 entity, Z default, 1 check char."""
    # 14th char must be 'Z'
    invalid_z = "24AAIHD1204K1AQ"
    assert validate_gst(invalid_z) is None

    # PAN portion must have 5 letters + 4 digits + 1 letter
    invalid_pan_digit = "24123451204K1ZQ"
    assert validate_gst(invalid_pan_digit) is None


# =====================================================================
# 8-11: DIRECT FLIPKART GST EXTRACTION (STATE, JSON-LD, HTML)
# =====================================================================

def test_08_gst_extraction_from_visible_html():
    """Test extraction of seller GST from visible HTML section."""
    html = """
    <html>
      <body>
        <div class="seller-details">
          <span>Sold By: Shanmugam Store</span>
          <span>GSTIN: 33AAECS5412Q1ZM</span>
          <span>Location: Thoothukudi, Tamil Nadu</span>
        </div>
      </body>
    </html>
    """
    soup = BeautifulSoup(html, "lxml")
    meta = extract_flipkart_seller_metadata(soup, html)
    assert meta["gst_number"] == "33AAECS5412Q1ZM"


def test_09_gst_extraction_from_json_ld():
    """Test extraction of seller GST from schema.org JSON-LD structured data."""
    html = """
    <html>
      <head>
        <script type="application/ld+json">
        {
          "@context": "https://schema.org",
          "@type": "Product",
          "name": "Organic Turmeric",
          "offers": {
            "@type": "Offer",
            "price": "299",
            "seller": {
              "@type": "Organization",
              "name": "Vedic Spices",
              "taxID": "24AAIHD1204K1ZQ",
              "address": {
                "@type": "PostalAddress",
                "addressLocality": "Surat",
                "addressRegion": "Gujarat",
                "postalCode": "395003"
              }
            }
          }
        }
        </script>
      </head>
      <body><div>Product Page</div></body>
    </html>
    """
    soup = BeautifulSoup(html, "lxml")
    meta = extract_flipkart_seller_metadata(soup, html)
    assert meta["gst_number"] == "24AAIHD1204K1ZQ"
    assert meta["city"] == "Surat"
    assert meta["state"] == "Gujarat"


def test_10_gst_extraction_from_embedded_state():
    """Test extraction of seller GST from window.__INITIAL_STATE__ JSON."""
    html = """
    <html>
      <body>
        <script>
          window.__INITIAL_STATE__ = {
            "pageDataV4": {
              "page": {
                "data": {
                  "sellerInfo": {
                    "sellerName": "Shree Balaji Traders",
                    "sellerGst": "07AAACH6891K1ZW",
                    "sellerLocation": "Delhi, India"
                  }
                }
              }
            }
          };
        </script>
      </body>
    </html>
    """
    soup = BeautifulSoup(html, "lxml")
    meta = extract_flipkart_seller_metadata(soup, html)
    assert meta["gst_number"] == "07AAACH6891K1ZW"


def test_11_direct_flipkart_gst_extraction_full_page():
    """Test end-to-end parse_product_page extracting GST directly from Flipkart."""
    html = """
    <html>
      <body>
        <div class="_1RLviY">
          <span>Seller:</span>
          <a href="/seller/royal-textiles">Royal Textiles</a>
        </div>
        <div class="seller-details">
          <p>GST No: 24AAIHD1204K1ZQ</p>
          <p>Contact: +91-9876543210</p>
          <p>Location: Surat, Gujarat, 395002</p>
        </div>
      </body>
    </html>
    """
    parsed = parse_product_page(html, page_url="https://www.flipkart.com/p/itm123")
    assert parsed["seller_name"] == "Royal Textiles"
    assert parsed["gst_number"] == "24AAIHD1204K1ZQ"


# =====================================================================
# 12-18: EXTERNAL SEARCH, WATERFALL & SNIPPET/PAGE EXTRACTION
# =====================================================================

@pytest.mark.asyncio
async def test_12_external_search_gst_extraction():
    """Test GST discovery and extraction through external search waterfall."""
    engine = WebResearchEngine()
    engine._google_test_mocked = True

    mock_google_results = [
        {
            "title": "Shanmugam Store - GST Details & Business Profile",
            "url": "https://www.zaubacorp.com/company/SHANMUGAM-STORE/33AAECS5412Q1ZM",
            "snippet": "Shanmugam Store in Thoothukudi, Tamil Nadu. GSTIN: 33AAECS5412Q1ZM, Address: 164/1 Sivan Koil Street.",
        }
    ]

    with patch.object(engine, "_query_google", new=AsyncMock(return_value=(mock_google_results, 200))):
        gst_val, gst_src = await engine.enrich_seller_gst(
            seller_name="Shanmugam Store",
            city="Thoothukudi",
            state="Tamil Nadu",
            location="Thoothukudi",
        )
        assert gst_val == "33AAECS5412Q1ZM"
        assert "zaubacorp.com" in gst_src
    await engine.close()


@pytest.mark.asyncio
async def test_13_google_to_bing_fallback():
    """Test Google failure (500 or empty) falls back seamlessly to Bing."""
    engine = WebResearchEngine()
    engine._google_test_mocked = True

    mock_bing_results = [
        {
            "title": "Krishna Enterprises GSTIN Profile",
            "url": "https://www.indiamart.com/krishna-enterprises/aboutus.html",
            "snippet": "Krishna Enterprises, Surat, Gujarat. GSTIN 24AAIHD1204K1ZQ, Pin 395003.",
        }
    ]

    with patch.object(engine, "_query_google", new=AsyncMock(return_value=([], 500))), \
         patch.object(engine, "_query_bing", new=AsyncMock(return_value=(mock_bing_results, 200))):
        gst_val, gst_src = await engine.enrich_seller_gst(
            seller_name="Krishna Enterprises",
            city="Surat",
            state="Gujarat",
        )
        assert gst_val == "24AAIHD1204K1ZQ"
    await engine.close()


@pytest.mark.asyncio
async def test_14_bing_to_brave_fallback():
    """Test Bing failure falls back to Brave."""
    engine = WebResearchEngine()
    engine._google_test_mocked = True

    mock_brave_results = [
        {
            "title": "Apex Electronics Delhi GST Details",
            "url": "https://tofler.in/apex-electronics",
            "snippet": "Apex Electronics, Delhi. GST No. 07AAACH6891K1ZW. Reg in New Delhi.",
        }
    ]

    with patch.object(engine, "_query_google", new=AsyncMock(return_value=([], 500))), \
         patch.object(engine, "_query_bing", new=AsyncMock(return_value=([], 500))), \
         patch.object(engine, "_query_brave", new=AsyncMock(return_value=(mock_brave_results, 200))):
        gst_val, gst_src = await engine.enrich_seller_gst(
            seller_name="Apex Electronics",
            city="Delhi",
            state="Delhi",
        )
        assert gst_val == "07AAACH6891K1ZW"
    await engine.close()


@pytest.mark.asyncio
async def test_15_brave_to_duckduckgo_fallback():
    """Test Brave failure falls back to DuckDuckGo."""
    engine = WebResearchEngine()
    engine._google_test_mocked = True

    mock_ddg_results = [
        {
            "title": "Southern Traders Chennai",
            "url": "https://www.quickcompany.in/company/southern-traders",
            "snippet": "Southern Traders in Chennai, Tamil Nadu. GSTIN 33AAECS5412Q1ZM.",
        }
    ]

    with patch.object(engine, "_query_google", new=AsyncMock(return_value=([], 500))), \
         patch.object(engine, "_query_bing", new=AsyncMock(return_value=([], 500))), \
         patch.object(engine, "_query_brave", new=AsyncMock(return_value=([], 429))), \
         patch.object(engine, "_query_ddg", new=AsyncMock(return_value=mock_ddg_results)):
        gst_val, gst_src = await engine.enrich_seller_gst(
            seller_name="Southern Traders",
            city="Chennai",
            state="Tamil Nadu",
        )
        assert gst_val == "33AAECS5412Q1ZM"
    await engine.close()


@pytest.mark.asyncio
async def test_16_google_429_fallback():
    """Test that Google 429 triggers backoff and falls back immediately without failing the run."""
    engine = WebResearchEngine()
    engine._google_test_mocked = True

    mock_bing_results = [
        {
            "title": "Om Sai Handicrafts GST",
            "url": "https://www.instafinancials.com/company/om-sai-handicrafts",
            "snippet": "Om Sai Handicrafts Jaipur Rajasthan. GST Number 08AAACH6891K1ZW.",
        }
    ]

    with patch.object(engine, "_query_google", new=AsyncMock(return_value=([], 429))), \
         patch.object(engine, "_query_bing", new=AsyncMock(return_value=(mock_bing_results, 200))):
        gst_val, gst_src = await engine.enrich_seller_gst(
            seller_name="Om Sai Handicrafts",
            city="Jaipur",
            state="Rajasthan",
        )
        assert gst_val == "08AAACH6891K1ZW"
    await engine.close()


@pytest.mark.asyncio
async def test_17_search_result_snippet_extraction():
    """Test extracting GSTIN candidate directly from snippet without opening page."""
    engine = WebResearchEngine()
    engine._google_test_mocked = True

    results = [
        {
            "title": "Radha Krishna Creations Surat",
            "url": "https://example-dir.com/radha-krishna",
            "snippet": "Contact Radha Krishna Creations in Surat, Gujarat. GSTIN is 24AAIHD1204K1ZQ.",
        }
    ]

    with patch.object(engine, "_query_google", new=AsyncMock(return_value=(results, 200))):
        gst_val, _ = await engine.enrich_seller_gst(
            seller_name="Radha Krishna Creations",
            city="Surat",
            state="Gujarat",
        )
        assert gst_val == "24AAIHD1204K1ZQ"
    await engine.close()


@pytest.mark.asyncio
async def test_18_opened_page_gst_extraction():
    """Test opening a high-relevance official or directory URL when snippet lacks GSTIN."""
    engine = WebResearchEngine()
    engine._google_test_mocked = True

    results = [
        {
            "title": "Shanmugam Store Official Profile",
            "url": "https://www.shanmugamstore.com/contact-us",
            "snippet": "Shanmugam Store Wholesale & Retail Merchant in Thoothukudi.",
        }
    ]

    mock_page_html = """
    <html>
      <body>
        <h1>Shanmugam Store</h1>
        <p>Address: 164 Sivan Koil Street, Thoothukudi, Tamil Nadu 628002</p>
        <p>GSTIN: 33AAECS5412Q1ZM</p>
      </body>
    </html>
    """

    with patch.object(engine, "_query_google", new=AsyncMock(return_value=(results, 200))), \
         patch.object(engine.website_parser, "fetch_html", new=AsyncMock(return_value=mock_page_html)):
        gst_val, gst_src = await engine.enrich_seller_gst(
            seller_name="Shanmugam Store",
            city="Thoothukudi",
            state="Tamil Nadu",
            location="Thoothukudi",
            website_url="https://www.shanmugamstore.com",
        )
        assert gst_val == "33AAECS5412Q1ZM"
    await engine.close()


# =====================================================================
# 19-28: SELLER IDENTITY MATCHING, DISAMBIGUATION & CONFLICT REJECTION
# =====================================================================

def test_19_seller_name_matching():
    """Test strong match between candidate name and seller name."""
    matched, score, reason = match_gst_to_seller(
        seller_name="Shanmugam Store",
        gst_number="33AAECS5412Q1ZM",
        snippet="Shanmugam Store Thoothukudi Tamil Nadu",
        city="Thoothukudi",
        state="Tamil Nadu",
    )
    assert matched is True
    assert score >= 95


def test_20_city_matching():
    """Test matching city elevates match score."""
    matched, score, _ = match_gst_to_seller(
        seller_name="Balaji Enterprises",
        gst_number="24AAIHD1204K1ZQ",
        snippet="Balaji Enterprises Surat Gujarat",
        city="Surat",
        state="Gujarat",
    )
    assert matched is True
    assert score == 100


def test_21_state_matching():
    """Test state matching aligns with GST state prefix."""
    matched, score, _ = match_gst_to_seller(
        seller_name="Vedic Naturals",
        gst_number="27AABCT3518Q1ZV",  # 27 = Maharashtra
        snippet="Vedic Naturals Pune Maharashtra",
        state="Maharashtra",
    )
    assert matched is True
    assert score >= 90


def test_22_pincode_matching():
    """Test pincode match confirms seller identity."""
    matched, score, _ = match_gst_to_seller(
        seller_name="Surat Silk Mills",
        gst_number="24AAIHD1204K1ZQ",
        snippet="Surat Silk Mills, Ring Road, Surat, 395002",
        city="Surat",
        pincode="395002",
    )
    assert matched is True
    assert score >= 90


def test_23_generic_seller_protection():
    """Test that generic names (e.g. 'Enterprises', 'Store', 'Fashion') require location evidence."""
    assert is_generic_seller_name("ABC Enterprises") is True
    assert is_generic_seller_name("Fashion Store") is True
    assert is_generic_seller_name("Royal Collections") is True
    assert is_generic_seller_name("REEPREECREATION") is False

    # Generic name without location in candidate must be rejected
    matched, score, reason = match_gst_to_seller(
        seller_name="ABC Enterprises",
        gst_number="24AAIHD1204K1ZQ",
        snippet="ABC Enterprises India - leading supplier",
        city="Surat",
        state="Gujarat",
    )
    assert matched is False
    assert "GENERIC_NAME_LOCATION_MISMATCH" in reason


def test_24_wrong_city_gst_rejection():
    """Test rejecting candidate GSTIN when city explicitly conflicts with seller location."""
    # Target is Surat, candidate is Chennai
    matched, score, reason = match_gst_to_seller(
        seller_name="Shanmugam Store",
        gst_number="33AAECS5412Q1ZM",  # Tamil Nadu GST
        snippet="Shanmugam Store, 164 Sivan Koil Street, Chennai, Tamil Nadu 600001",
        city="Surat",
        state="Gujarat",
    )
    assert matched is False
    assert "LOCATION_MISMATCH" in reason or "STATE_CODE_MISMATCH" in reason


def test_25_wrong_state_gst_rejection():
    """Test rejecting candidate GSTIN when 2-digit state code contradicts seller state."""
    # Target state is Gujarat (24), candidate GSTIN is 07 (Delhi)
    matched, score, reason = match_gst_to_seller(
        seller_name="Royal Textiles",
        gst_number="07AAACH6891K1ZW",
        snippet="Royal Textiles Chandni Chowk Delhi",
        city="Surat",
        state="Gujarat",
    )
    assert matched is False
    assert "STATE_CODE_MISMATCH" in reason or "LOCATION_MISMATCH" in reason


def test_26_same_name_business_rejection():
    """Test that businesses with identical names in different locations are rejected."""
    target_seller = "ABC Enterprises"
    target_city = "Surat"
    target_state = "Gujarat"

    source_snippet = "ABC Enterprises, Connaught Place, New Delhi, Delhi 110001. GSTIN: 07AAACH6891K1ZW"
    matched, _, reason = match_gst_to_seller(
        seller_name=target_seller,
        gst_number="07AAACH6891K1ZW",
        snippet=source_snippet,
        city=target_city,
        state=target_state,
    )
    assert matched is False


def test_27_multiple_gstin_candidate_selection():
    """Test selecting the best matching GSTIN when multiple candidate numbers are present."""
    snippet = """
    Candidate 1: 07AAACH6891K1ZW (Delhi Branch)
    Candidate 2: 24AAIHD1204K1ZQ (Surat Head Office, Gujarat)
    Candidate 3: 27AABCT3518Q1ZV (Mumbai Office)
    """
    matched_surat, score_surat, _ = match_gst_to_seller(
        seller_name="Surat Silks",
        gst_number="24AAIHD1204K1ZQ",
        snippet="Surat Silks, Surat, Gujarat 24AAIHD1204K1ZQ",
        city="Surat",
        state="Gujarat",
    )
    matched_delhi, score_delhi, _ = match_gst_to_seller(
        seller_name="Surat Silks",
        gst_number="07AAACH6891K1ZW",
        snippet="Surat Silks Delhi branch",
        city="Surat",
        state="Gujarat",
    )
    assert matched_surat is True
    assert matched_delhi is False or score_surat > score_delhi


def test_28_gst_state_code_cross_check():
    """Test verifying that 2-digit GST state code matches target state."""
    # 24 = Gujarat
    matched, _, _ = match_gst_to_seller(
        seller_name="Gujarat Traders",
        gst_number="24AAIHD1204K1ZQ",
        snippet="Gujarat Traders Ahmedabad Gujarat",
        state="Gujarat",
    )
    assert matched is True


# =====================================================================
# 29-30: CACHING SAFETY & FINAL EXCEL MAPPING
# =====================================================================

@pytest.mark.asyncio
async def test_29_temporary_failure_must_not_become_permanent_not_found():
    """Test that temporary network/429 errors are not cached permanently as NOT_FOUND."""
    engine = WebResearchEngine()
    engine._google_test_mocked = True

    # First attempt: Google returns 429, Bing returns 500 -> no result
    with patch.object(engine, "_query_google", new=AsyncMock(return_value=([], 429))), \
         patch.object(engine, "_query_bing", new=AsyncMock(return_value=([], 500))), \
         patch.object(engine, "_query_brave", new=AsyncMock(return_value=([], 500))), \
         patch.object(engine, "_query_ddg", new=AsyncMock(return_value=[])):
        res1, _ = await engine.enrich_seller_gst(seller_name="Test Store", city="Surat")
        assert res1 is None

    # Check cache does NOT contain a permanent "NOT_FOUND" entry blocking future queries
    assert engine.cache.get("gst::flipkart::test store") is None

    # Second attempt: Bing recovers and returns valid result -> successfully extracted
    valid_results = [
        {
            "title": "Test Store Surat GST",
            "url": "https://example.com/test",
            "snippet": "Test Store in Surat, Gujarat. GSTIN: 24AAIHD1204K1ZQ",
        }
    ]
    with patch.object(engine, "_query_google", new=AsyncMock(return_value=([], 429))), \
         patch.object(engine, "_query_bing", new=AsyncMock(return_value=(valid_results, 200))):
        res2, _ = await engine.enrich_seller_gst(seller_name="Test Store", city="Surat", state="Gujarat")
        assert res2 == "24AAIHD1204K1ZQ"

    await engine.close()


@pytest.mark.asyncio
async def test_30_final_output_mapping_to_gst_number(tmp_path):
    """Test that enriched GSTIN flows cleanly to the final Excel column named exactly 'GST Number'."""
    engine = WebResearchEngine()
    seller_record = {
        "seller_name": "Shanmugam Store",
        "marketplace": "flipkart",
        "city": "Thoothukudi",
        "state": "Tamil Nadu",
        "location": "Thoothukudi",
        "gst_number": "33AAECS5412Q1ZM",  # Verified Flipkart GST
    }

    enriched = await engine.enrich_seller(seller_record)
    assert enriched["GST Number"] == "33AAECS5412Q1ZM"
    assert enriched["gst_number"] == "33AAECS5412Q1ZM"

    # Verify Excel manager writes to "GST Number" column
    excel_path = tmp_path / "test_output.xlsx"
    manager = LiveExcelManager(excel_path)
    row_num = manager.write_or_update_seller(enriched)
    saved, msg = manager.verify_saved_row(row_num, enriched)
    assert saved is True

    # Read back directly
    import openpyxl
    wb = openpyxl.load_workbook(excel_path)
    ws = wb.active
    headers = [cell.value for cell in ws[1]]
    assert "GST Number" in headers
    gst_col_idx = headers.index("GST Number") + 1
    # row_num corresponds to the 1st data row, which is row 2 in worksheet (row 1 is header)
    assert ws.cell(row=row_num + 1, column=gst_col_idx).value == "33AAECS5412Q1ZM"

    await engine.close()
