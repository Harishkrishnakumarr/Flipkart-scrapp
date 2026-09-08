"""Comprehensive unit tests for PAN Number extraction from verified GSTIN."""

import pytest
from unittest.mock import AsyncMock, patch
import openpyxl

from scraper.exporter import LiveExcelManager
from scraper.validator import (
    extract_pan_from_gstin,
    validate_gst,
    validate_pan,
)
from scraper.web_research import WebResearchEngine


# =====================================================================
# 1-5: GSTIN -> PAN EXTRACTION & NORMALIZATION
# =====================================================================

def test_01_valid_gstin_to_correct_pan():
    """Test standard 15-char GSTIN extracts exact 10-char PAN."""
    gstin = "24AAIHD1204K1ZQ"
    pan = extract_pan_from_gstin(gstin)
    assert pan == "AAIHD1204K"
    assert len(pan) == 10


def test_02_another_valid_gstin_to_correct_pan():
    """Test a second valid GSTIN with Delhi state code extracts exact PAN."""
    gstin = "07AAGPA0669R1ZD"
    pan = extract_pan_from_gstin(gstin)
    assert pan == "AAGPA0669R"
    assert len(pan) == 10


def test_03_lowercase_gstin_normalization():
    """Test lowercase GSTIN string is normalized and uppercase PAN is returned."""
    gstin = "24aaihd1204k1zq"
    pan = extract_pan_from_gstin(gstin)
    assert pan == "AAIHD1204K"


def test_04_whitespace_normalization():
    """Test surrounding and internal whitespace in GSTIN string is stripped."""
    gstin = "  24AAIHD1204K1ZQ  "
    pan = extract_pan_from_gstin(gstin)
    assert pan == "AAIHD1204K"

    gstin_spaced = "24 aaihd 1204 k 1 zq"
    pan_spaced = extract_pan_from_gstin(gstin_spaced)
    assert pan_spaced == "AAIHD1204K"


def test_05_formatted_gstin_normalization():
    """Test hyphen-separated GSTIN is cleanly normalized to 10-char PAN."""
    gstin = "24-AAIHD-1204-K-1ZQ"
    pan = extract_pan_from_gstin(gstin)
    assert pan == "AAIHD1204K"


# =====================================================================
# 6-10: FORMAT STRUCTURE & BOUNDARY VALIDATION
# =====================================================================

def test_06_invalid_gstin_rejected():
    """Test malformed or invalid GSTINs do not extract a PAN."""
    assert extract_pan_from_gstin("INVALID_GSTIN") is None
    assert extract_pan_from_gstin("99AAIHD1204K1ZQ") is None  # Invalid state code 99
    assert extract_pan_from_gstin("24AAIHD1204K") is None      # Too short (12 chars)
    assert extract_pan_from_gstin("") is None
    assert extract_pan_from_gstin(None) is None


def test_07_invalid_pan_structure_rejected():
    """Test invalid PAN patterns (wrong length, bad 4th char entity type) are rejected."""
    # 4th char 'X' is not a valid Indian PAN entity type (C, P, H, F, A, T, B, L, J, G)
    assert validate_pan("AAIXD1204K") is None
    assert validate_pan("12345ABCDE") is None
    assert validate_pan("ABCDE12345") is None
    assert validate_pan("ABCDE1234") is None   # 9 chars
    assert validate_pan("ABCDE12345F") is None # 11 chars
    assert validate_pan("ABC@D1234F") is None


def test_08_gst_state_code_excluded_from_pan():
    """Test 2-digit state code is excluded from extracted PAN."""
    gstin = "24AAIHD1204K1ZQ"
    pan = extract_pan_from_gstin(gstin)
    assert not pan.startswith("24")
    assert pan == "AAIHD1204K"


def test_09_entity_number_excluded_from_pan():
    """Test 13th entity number character is excluded from extracted PAN."""
    gstin = "24AAIHD1204K1ZQ"
    pan = extract_pan_from_gstin(gstin)
    # The 13th char is '1' in 24AAIHD1204K1ZQ
    assert not pan.endswith("1")
    assert pan == "AAIHD1204K"


def test_10_checksum_character_excluded_from_pan():
    """Test 15th checksum character is excluded from extracted PAN."""
    gstin = "24AAIHD1204K1ZQ"
    pan = extract_pan_from_gstin(gstin)
    # The 15th char is 'Q' in 24AAIHD1204K1ZQ
    assert pan != "AAIHD1204Q"
    assert pan == "AAIHD1204K"


# =====================================================================
# 11-13: EMBEDDED PAN CONSISTENCY & CROSS-CHECKS
# =====================================================================

def test_11_gst_derived_pan_matches_gst_embedded_pan():
    """Test GST-derived PAN strictly equals gstin[2:12]."""
    gstin = "33AAECS5412Q1ZM"
    pan = extract_pan_from_gstin(gstin)
    assert pan == gstin[2:12]
    assert pan == "AAECS5412Q"


@pytest.mark.asyncio
async def test_12_existing_pan_match():
    """Test matching existing PAN is confirmed and preserved."""
    engine = WebResearchEngine()
    seller_record = {
        "seller_name": "Shanmugam Store",
        "marketplace": "flipkart",
        "city": "Thoothukudi",
        "state": "Tamil Nadu",
        "gst_number": "33AAECS5412Q1ZM",
        "pan_number": "AAECS5412Q",  # Matching PAN
    }
    with patch.object(engine, "_query_google", new=AsyncMock(return_value=([], 200))), \
         patch.object(engine, "_query_bing", new=AsyncMock(return_value=([], 200))), \
         patch.object(engine, "_query_brave", new=AsyncMock(return_value=([], 200))), \
         patch.object(engine, "_query_ddg", new=AsyncMock(return_value=[])):
        enriched = await engine.enrich_seller(seller_record)
    assert enriched["PAN Number"] == "AAECS5412Q"
    assert enriched["pan_number"] == "AAECS5412Q"
    await engine.close()


@pytest.mark.asyncio
async def test_13_existing_pan_mismatch_detection(caplog):
    """Test conflicting existing PAN logs warning and GST-derived PAN is authoritative."""
    import logging
    engine = WebResearchEngine()
    seller_record = {
        "seller_name": "Shanmugam Store",
        "marketplace": "flipkart",
        "city": "Thoothukudi",
        "state": "Tamil Nadu",
        "gst_number": "33AAECS5412Q1ZM",  # Embedded PAN: AAECS5412Q
        "pan_number": "BBBBB1234C",        # Conflicting existing PAN
    }
    with patch.object(engine, "_query_google", new=AsyncMock(return_value=([], 200))), \
         patch.object(engine, "_query_bing", new=AsyncMock(return_value=([], 200))), \
         patch.object(engine, "_query_brave", new=AsyncMock(return_value=([], 200))), \
         patch.object(engine, "_query_ddg", new=AsyncMock(return_value=[])), \
         caplog.at_level(logging.WARNING):
        enriched = await engine.enrich_seller(seller_record)
    
    assert enriched["PAN Number"] == "AAECS5412Q"
    assert enriched["pan_number"] == "AAECS5412Q"
    assert "PAN MISMATCH" in caplog.text or "BBBBB1234C" in caplog.text
    await engine.close()


# =====================================================================
# 14-16: NO VERIFIED GST, OUTPUT EXCEL MAPPING & MULTIPLE GSTINs
# =====================================================================

@pytest.mark.asyncio
async def test_14_no_verified_gst_yields_no_pan():
    """Test that absence of verified GSTIN results in blank/None PAN without fabrication."""
    engine = WebResearchEngine()
    seller_record = {
        "seller_name": "Unverified Store",
        "marketplace": "flipkart",
        "city": "Surat",
        "state": "Gujarat",
        "gst_number": None,
    }

    with patch.object(engine, "_query_google", new=AsyncMock(return_value=([], 200))), \
         patch.object(engine, "_query_bing", new=AsyncMock(return_value=([], 200))), \
         patch.object(engine, "_query_brave", new=AsyncMock(return_value=([], 200))), \
         patch.object(engine, "_query_ddg", new=AsyncMock(return_value=[])):
        enriched = await engine.enrich_seller(seller_record)

    assert enriched.get("PAN Number") is None
    assert enriched.get("pan_number") is None
    await engine.close()


@pytest.mark.asyncio
async def test_15_final_output_mapping_to_pan_number(tmp_path):
    """Test that derived PAN is written to exact Excel column named 'PAN Number'."""
    engine = WebResearchEngine()
    seller_record = {
        "seller_name": "Shanmugam Store",
        "marketplace": "flipkart",
        "city": "Thoothukudi",
        "state": "Tamil Nadu",
        "location": "Thoothukudi",
        "gst_number": "33AAECS5412Q1ZM",
    }

    with patch.object(engine, "_query_google", new=AsyncMock(return_value=([], 200))), \
         patch.object(engine, "_query_bing", new=AsyncMock(return_value=([], 200))), \
         patch.object(engine, "_query_brave", new=AsyncMock(return_value=([], 200))), \
         patch.object(engine, "_query_ddg", new=AsyncMock(return_value=[])):
        enriched = await engine.enrich_seller(seller_record)
    assert enriched["PAN Number"] == "AAECS5412Q"

    excel_path = tmp_path / "test_pan_output.xlsx"
    manager = LiveExcelManager(excel_path)
    row_num = manager.write_or_update_seller(enriched)
    saved, msg = manager.verify_saved_row(row_num, enriched)
    assert saved is True

    wb = openpyxl.load_workbook(excel_path)
    ws = wb.active
    headers = [cell.value for cell in ws[1]]
    assert "PAN Number" in headers
    pan_col_idx = headers.index("PAN Number") + 1
    # row_num + 1 corresponds to data row in worksheet (row 1 is header)
    assert ws.cell(row=row_num + 1, column=pan_col_idx).value == "AAECS5412Q"
    await engine.close()


def test_16_multiple_gstin_handling():
    """Test extracting embedded PAN from candidate GSTINs and verifying consistency."""
    gstin_1 = "24AAIHD1204K1ZQ"  # Gujarat registration
    gstin_2 = "27AAIHD1204K1ZR"  # Maharashtra registration (same PAN)

    pan_1 = extract_pan_from_gstin(gstin_1)
    pan_2 = extract_pan_from_gstin(gstin_2)

    assert pan_1 == "AAIHD1204K"
    assert pan_2 == "AAIHD1204K"
    assert pan_1 == pan_2
